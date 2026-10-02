"""税局财务报表导入文件（v0.34）：导出 → 电子税务局「财务报表报送」导入 → 申报后把实际提交的文件导回登记。

2026-10-01 与用户定：
- 流程：月结锁期 → 导出本文件 → 税局导入并申报 → 把**实际提交的那份**（网页上改过数就先改文件）导回 ERP 登记，
  原件存数据目录 tax-filing/{期间}/。税局申报后只给 PDF/查询页，没有导入格式的文件，所以回填用提交的文件。
- 季报只报资产负债表 + 利润表（按备案带出，2026-Q2 税局页面确认）；模板里的现金流量表留 0，年报前再做。
- 模板 = 2026-Q1 导入成功的那份 xlsx 去掉数（app/templates/fs_small_quarter.xlsx）；全部写值不写公式
  （openpyxl 写的公式没有缓存值，导入方读到的是空）。税局导入接受 xls、xlsx，导回两种都认。
- 资产负债表：期末 = ERP；年初 = 本年已登记申报的年初数（照抄已申报，用户定），没有才用 ERP 上年末。
- 利润表：本年累计 = ERP；本期 = ERP 本年累计 − 上期已申报累计（statements.statements 的申报口径）。
"""
import base64
import io
import re
from datetime import date, datetime
from pathlib import Path

from fastapi import Depends, HTTPException
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel, ConfigDict, Field
from sqlalchemy import select

from . import operations as O
from .paths import data_dir

TEMPLATE = Path(__file__).parent / "templates" / "fs_small_quarter.xlsx"
IS_KEYS = {1: "revenue", 2: "cost", 3: "taxes", 11: "selling", 14: "admin", 18: "finance", 22: "nonop_income",
           24: "nonop_expense", 30: "profit_total", 31: "income_tax", 32: "net_profit"}
BS_KEYS = {30: "total_assets", 47: "total_liabilities", 52: "equity"}


def filing_dir() -> Path:
    return data_dir() / "tax-filing"


def quarter_dates(period: str) -> tuple[date, date]:
    y, q = int(period[:4]), int(period[-1])
    end = date(y + (q == 4), q * 3 % 12 + 1, 1)
    return date(y, q * 3 - 2, 1), date.fromordinal(end.toordinal() - 1)


# ---------- 导出 ----------
def export(db, period: str) -> bytes:
    import openpyxl
    from .company import get as company
    from .statements import statements
    if not re.fullmatch(r"\d{4}-Q[1-4]", period):
        raise HTTPException(400, "税局导入文件按季度导出（期间如 2026-Q3）；年报模板另做")
    d = statements(db, period)
    c = company(db)
    start, end = quarter_dates(period)
    wb = openpyxl.load_workbook(TEMPLATE)
    head = [c.get("tax_no", ""), c.get("name", ""), start.isoformat(), end.isoformat()]
    for ws in wb:
        cells = ("D3", "H3", "D4", "H4") if ws.title.startswith("资产负债表") else ("D3", "F3", "D4", "F4")
        for addr, v in zip(cells, head):
            ws[addr].value = v
    bs = {x["row"]: x for x in d["balance_sheet"]}
    is_ = {x["row"]: x for x in d["income_statement"]}
    for ws in wb:
        kind = _kind(ws.title)
        for row in ws.iter_rows(min_row=6):
            for cell in row:
                if not (isinstance(cell.value, int) and cell.column_letter in ("C", "G")):
                    continue
                r = cell.value
                if kind == "bs" and r in bs:
                    ws.cell(cell.row, cell.column + 1).value = bs[r]["amount"]
                    ws.cell(cell.row, cell.column + 2).value = bs[r]["begin"]
                elif kind == "is" and r in is_:
                    ws.cell(cell.row, cell.column + 1).value = is_[r]["period_filed"]
                    ws.cell(cell.row, cell.column + 2).value = is_[r]["ytd"]
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------- 导回 ----------
def _kind(title: str) -> str:
    return "bs" if "资产负债" in title else "is" if "利润" in title else ""


def _grid(name: str, data: bytes) -> dict[str, list[list]]:
    """{表名: 行列值}，xls 与 xlsx 统一成二维列表（0 基）"""
    if name.lower().endswith(".xls"):
        import xlrd
        wb = xlrd.open_workbook(file_contents=data)
        out = {}
        for sh in wb.sheets():
            rows = []
            for r in range(sh.nrows):
                vals = []
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    v = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        v = xlrd.xldate.xldate_as_datetime(v, wb.datemode)
                    vals.append(v)
                rows.append(vals)
            out[sh.name] = rows
        return out
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    return {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb}


def _num(v) -> float:
    if v in (None, ""):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(",", "").replace("，", "").strip()
    try:
        return float(s) if s not in ("", "-", "—") else 0.0
    except ValueError:
        raise HTTPException(400, f"金额读不出来：{v}")


def _date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    m = re.search(r"(\d{4})[-/年.](\d{1,2})[-/月.](\d{1,2})", str(v or ""))
    return date(int(m[1]), int(m[2]), int(m[3])) if m else None


def _after(rows, label):
    """表头「标签 → 右边第一个非空值」"""
    for r in rows[:8]:
        for i, v in enumerate(r):
            if isinstance(v, str) and v.strip() == label:
                for w in r[i + 1:]:
                    if w not in (None, ""):
                        return w
    return None


def parse(name: str, data: bytes) -> dict:
    sheets = _grid(name, data)
    bs, is_, head = {}, {}, None
    for title, rows in sheets.items():
        kind = _kind(title)
        if not kind:
            continue
        head = head or rows
        hdr_i = next((i for i, r in enumerate(rows[:10]) if any(isinstance(v, str) and v.strip() == "行次" for v in r)), None)
        if hdr_i is None:
            raise HTTPException(400, f"「{title}」找不到行次表头，不是税局财务报表格式")
        hdr = rows[hdr_i]
        cols = [i for i, v in enumerate(hdr) if isinstance(v, str) and v.strip() == "行次"]
        for ci in cols:
            # 行次右边两列的表头决定含义：资产负债表 期末/年初；利润表 本期/本年累计（年报是 本年累计/上年）
            labels = [str(hdr[ci + k] or "") for k in (1, 2)] if ci + 2 < len(hdr) else ["", ""]
            if kind == "is":
                ytd_k = next((k for k, l in zip((1, 2), labels) if "本年累计" in l), 2)
                cur_k = next((k for k, l in zip((1, 2), labels) if "本期" in l), None)
            for r in rows[hdr_i + 1:]:
                if ci >= len(r) or not isinstance(r[ci], (int, float)) or isinstance(r[ci], bool):
                    continue
                n = int(r[ci])
                get = lambda k: _num(r[ci + k]) if ci + k < len(r) else 0.0
                if kind == "bs":
                    bs[n] = [round(get(1), 2), round(get(2), 2)]
                else:
                    is_[n] = [round(get(cur_k), 2) if cur_k else None, round(get(ytd_k), 2)]
    if not bs or not is_:
        raise HTTPException(400, "文件里要有资产负债表和利润表两张表（税局「财务报表报送与信息采集」导入格式）")
    start, end = _date(_after(head, "所属期起")), _date(_after(head, "所属期止"))
    if not (start and end):
        raise HTTPException(400, "表头读不到所属期起止")
    if start.month == 1 and end.month == 12:
        period = str(end.year)
    elif end.month % 3 == 0:
        period = f"{end.year}-Q{end.month // 3}"
    else:
        raise HTTPException(400, f"所属期 {start}~{end} 不是季度或年度")
    return {"period": period, "start": start.isoformat(), "end": end.isoformat(),
            "tax_no": str(_after(head, "纳税人识别号") or "").strip(), "bs": bs, "is": is_}


def register(db, name: str, data: bytes, filed_date=None, remark: str = "") -> dict:
    from .company import get as company
    from .tax_filing import TaxFiling
    p = parse(name, data)
    own = company(db).get("tax_no", "")
    if own and p["tax_no"] and p["tax_no"] != own:
        raise HTTPException(400, f"纳税人识别号 {p['tax_no']} 不是本单位（{own}）")
    # 与 ERP 导出数对比：看出用户在税局网页上改了什么
    diffs = []
    if "-Q" in p["period"]:
        from .statements import statements
        d = statements(db, p["period"])
        mine = {("bs", x["row"]): (x["amount"], x["begin"]) for x in d["balance_sheet"]}
        mine |= {("is", x["row"]): (x["period_filed"], x["ytd"]) for x in d["income_statement"]}
        labels = {("bs", x["row"]): x["label"] for x in d["balance_sheet"]} | {("is", x["row"]): x["label"] for x in d["income_statement"]}
        for kind in ("bs", "is"):
            for n, vals in sorted(p[kind].items()):
                erp = mine.get((kind, n))
                if not erp:
                    continue
                for col, a, b in zip(("期末" if kind == "bs" else "本期", "年初" if kind == "bs" else "本年累计"), vals, erp):
                    if a is not None and abs(a - b) > 0.005:
                        diffs.append({"sheet": "资产负债表" if kind == "bs" else "利润表", "row": n, "label": labels[(kind, n)],
                                      "col": col, "filed": a, "erp": b})
    t = db.scalar(select(TaxFiling).where(TaxFiling.period == p["period"]))
    if not t:
        t = TaxFiling(period=p["period"])
        db.add(t)
    vals = {k: p["is"].get(n, [0, 0])[1] for n, k in IS_KEYS.items()}
    vals |= {k: p["bs"].get(n, [0, 0])[0] for n, k in BS_KEYS.items()}
    folder = filing_dir() / p["period"]
    folder.mkdir(parents=True, exist_ok=True)
    safe = re.sub(r'[\\/:*?"<>|]', "_", Path(name).name)
    (folder / safe).write_bytes(data)
    files = sorted(set((t.data or {}).get("files", [])) | {f"{p['period']}/{safe}"})
    t.data = {**vals, "bs": {str(k): v for k, v in p["bs"].items()}, "is": {str(k): v for k, v in p["is"].items()}, "files": files}
    t.filed_date = filed_date or t.filed_date or date.today()
    t.remark = remark or t.remark or f"导入已申报文件 {safe}"
    O.audit(db, "导入已申报报表", period=p["period"], file=safe, diffs=len(diffs))
    db.commit()
    return {"period": p["period"], "start": p["start"], "end": p["end"], "file": safe, "diffs": diffs}


class ImportIn(BaseModel):
    model_config = ConfigDict(str_strip_whitespace=True)
    name: str = Field(min_length=1)
    content_b64: str = Field(min_length=1)
    filed_date: date | None = None
    remark: str = ""


def install(app, get_db):
    from urllib.parse import quote

    @app.get("/api/tax-filings/export.xlsx")
    def export_api(period: str, db=Depends(get_db)):
        data = export(db, period)
        name = quote(f"财务报表报送与信息采集（小企业会计准则）{period}.xlsx")
        return StreamingResponse(io.BytesIO(data), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                 headers={"Content-Disposition": f"attachment; filename*=utf-8''{name}"})

    @app.post("/api/tax-filings/import")
    def import_api(data: ImportIn, db=Depends(get_db)):
        if not data.name.lower().endswith((".xls", ".xlsx")):
            raise HTTPException(400, "只认 xls / xlsx（税局财务报表导入格式）")
        try:
            raw = base64.b64decode(data.content_b64)
        except Exception:
            raise HTTPException(400, "文件内容不是 base64")
        return register(db, data.name, raw, data.filed_date, data.remark)

    @app.get("/api/tax-filings/file")
    def file_api(path: str):
        f = (filing_dir() / path).resolve()
        if filing_dir().resolve() not in f.parents or not f.is_file():
            raise HTTPException(404, "文件不存在")
        return FileResponse(f, filename=f.name)
