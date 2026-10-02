"""开票导出（v0.32）：销售合同 → 税局「批量开票-导入开票模板」（数电，V260401 版，模板在 app/templates/）。

填两张表：1-发票基本信息（发票流水号、类型、含税、购方名称/税号/地址/电话/开户行/账号、备注）、
2-发票明细信息（项目名称、税收编码、规格型号、单位、数量、含税单价、金额、税率）。模板说明要求「所有内容按文本输入」。
取数规矩（2026-09-29，按历史开出的票归纳）：
- 项目名称 = 物料「开票名称」，空则用物料名称（税局按编码自动加「*制冷空调设备*」前缀）
- 税收编码 = 物料「税收编码」，空则用本单位设置的默认编码（1090131050000000000 制冷空调设备）
- 规格型号 = 这家客户的对方料号（物料别名 alias_spec，客户看得懂的写法）；「料号 xxx」这种内部编码不用，退回我方规格
- 含税单价取合同单价，金额 = 数量 × 单价（两位小数）；税率默认 13%
开出后下载发票（PDF/XML）到发票档案导入，按票号自动挂回合同。
"""
import io
import re
from pathlib import Path
from typing import Optional

from fastapi import Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, ConfigDict, Field
from sqlalchemy import inspect, select, text

from .models import Contract, ContractLine, Material, MaterialAlias

TEMPLATE = Path(__file__).parent / "templates" / "batch_invoice_V260401.xlsx"
DEFAULT_TAX_CODE = "1090131050000000000"


class ExportLine(BaseModel):
    model_config = ConfigDict(str_strip_whitespace=True)
    line_id: int
    qty: float = Field(gt=0)
    item_name: Optional[str] = None       # 覆盖项目名称
    spec: Optional[str] = None            # 覆盖规格型号
    tax_rate: float = Field(default=0.13, ge=0, lt=1)


class ExportIn(BaseModel):
    model_config = ConfigDict(str_strip_whitespace=True)
    serial: str = Field(default="", max_length=20)
    invoice_type: str = Field(default="增值税专用发票", pattern=r"^(增值税专用发票|普通发票)$")
    remark: str = Field(default="", max_length=230)
    lines: list[ExportLine] = Field(min_length=1)


def ensure_schema(engine):
    cols = {c["name"] for c in inspect(engine).get_columns("material")}
    with engine.begin() as conn:
        if "invoice_name" not in cols:
            conn.execute(text("ALTER TABLE material ADD COLUMN invoice_name VARCHAR(100) NOT NULL DEFAULT ''"))
        if "tax_code" not in cols:
            conn.execute(text("ALTER TABLE material ADD COLUMN tax_code VARCHAR(19) NOT NULL DEFAULT ''"))


def default_serial(contract_no: str) -> str:
    s = re.sub(r"[^0-9A-Za-z_-]", "", contract_no)   # 流水号限 20 字符，去掉中文（如「XS2605003-新」）
    return s[:20]


def spec_for(db, m: Material, partner_id: int) -> str:
    a = db.scalar(select(MaterialAlias).where(MaterialAlias.material_id == m.id, MaterialAlias.partner_id == partner_id))
    if a and a.alias_spec and not a.alias_spec.startswith("料号"):
        return a.alias_spec
    return m.spec or ""


def preview_lines(db, c: Contract) -> list[dict]:
    from .company import get
    code = get(db).get("default_tax_code") or DEFAULT_TAX_CODE
    out = []
    for l in c.lines:
        m = l.material
        if not l.qty or not l.price_tax:
            continue
        out.append({"line_id": l.id, "material_code": m.code, "material_name": m.name,
                    "item_name": m.invoice_name or m.name, "tax_code": m.tax_code or code,
                    "spec": spec_for(db, m, c.partner_id), "unit": m.unit, "qty": l.qty, "price_tax": l.price_tax,
                    "amount_tax": round(l.qty * l.price_tax, 2), "tax_rate": 0.13})
    return out


def build(db, c: Contract, data: ExportIn) -> bytes:
    import openpyxl
    from .company import get
    if c.contract_type != "sales":
        raise ValueError("只有销售合同开票")
    p = c.partner
    missing = [k for k, v in (("税号", p.tax_no), ("名称", p.name)) if not v]
    if data.invoice_type == "增值税专用发票" and missing:
        raise ValueError(f"开专票需要购方{'、'.join(missing)}，先在往来单位补齐")
    code = get(db).get("default_tax_code") or DEFAULT_TAX_CODE
    lines = {l.id: l for l in c.lines}
    serial = data.serial or default_serial(c.contract_no)
    wb = openpyxl.load_workbook(TEMPLATE)
    s1, s2 = wb["1-发票基本信息"], wb["2-发票明细信息"]
    head = {s1.cell(3, i).value: i for i in range(1, s1.max_column + 1) if s1.cell(3, i).value}
    row = {"发票流水号": serial, "发票类型": data.invoice_type, "是否含税": "是", "受票方自然人标识": "否",
           "购买方名称": p.name, "购买方纳税人识别号": p.tax_no or "", "购买方地址": p.address or "",
           "购买方电话": p.phone or "", "购买方开户银行": p.bank_name or "", "购买方银行账号": (p.bank_account or "").replace(" ", ""),
           "备注": data.remark or f"合同{c.contract_no}"}
    for k, v in row.items():
        if k in head and v:
            s1.cell(4, head[k], str(v))
    h2 = {s2.cell(3, i).value: i for i in range(1, s2.max_column + 1) if s2.cell(3, i).value}
    r = 4
    for x in data.lines:
        l = lines.get(x.line_id)
        if not l:
            raise ValueError("明细不属于该合同")
        m = l.material
        amount = round(x.qty * l.price_tax, 2)
        vals = {"发票流水号": serial, "项目名称": x.item_name or m.invoice_name or m.name, "商品和服务税收编码": m.tax_code or code,
                "规格型号": x.spec if x.spec is not None else spec_for(db, m, c.partner_id), "单位": m.unit,
                "数量": f"{x.qty:g}", "单价": f"{l.price_tax:g}", "金额": f"{amount:.2f}", "税率": f"{x.tax_rate:g}"}
        for k, v in vals.items():
            if k in h2 and v != "":
                s2.cell(r, h2[k], v)
        r += 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def install(app, get_db):
    @app.get("/api/contracts/{cid}/invoice-export")
    def export_preview(cid: int, db=Depends(get_db)):
        c = db.get(Contract, cid)
        if not c:
            raise HTTPException(404, "合同不存在")
        p = c.partner
        return {"serial": default_serial(c.contract_no), "remark": f"合同{c.contract_no}", "lines": preview_lines(db, c),
                "buyer": {"name": p.name, "tax_no": p.tax_no, "address": p.address, "phone": p.phone,
                          "bank_name": p.bank_name, "bank_account": p.bank_account}}

    @app.post("/api/contracts/{cid}/invoice-export")
    def export(cid: int, data: ExportIn, db=Depends(get_db)):
        from urllib.parse import quote
        c = db.get(Contract, cid)
        if not c:
            raise HTTPException(404, "合同不存在")
        raw = build(db, c, data)
        name = quote(f"批量开票-{c.contract_no}.xlsx")
        return StreamingResponse(io.BytesIO(raw), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                 headers={"Content-Disposition": f"attachment; filename*=utf-8''{name}"})
