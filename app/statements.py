"""小企业会计准则正式报表（v0.28）：资产负债表（会小企01表）、利润表（会小企02表）。

行次、项目名称与税务局「财务报表报送与信息采集（小企业会计准则）」表样一致，按科目 → 报表项目映射取数。
取数规则（《小企业会计准则》附录报表编制说明）：
- 往来按明细（往来单位）重分类：应收账款 = 应收借方 + 预收借方；预收账款 = 预收贷方 + 应收贷方；
  预付账款 = 预付借方 + 应付借方 + 暂估借方；应付账款 = 应付贷方 + 预付贷方 + 1234 货到票未到（暂估）贷方
- 增值税借方余额（待认证进项 222104、应交增值税明细净借方＝留抵）按财会〔2016〕22 号列「其他流动资产」，
  应交税费只反映欠税局的（2026-10-01 用户定；v0.28 曾按附录「借方以−号填列」，负债合计会变负数）
- 未分配利润 = 本年利润 4103 + 利润分配 4104 + 期初权益 4199 + 尚未结转的损益（表结法，月度不结转）
科目编码用企业会计准则编码（与原报税底稿一致），报表按小企业会计准则格式；换准则只加一套行定义
利润表「本期金额」给两个口径：账面（按期间发生额）与申报口径（本年累计 − 上期已申报累计，
tax_filing 登记的数；与用户定的「以前季度差异放在本期」一致）。
换准则时：新增一套 BS/IS 行定义即可，科目与凭证不动。
"""
import io
from collections import defaultdict
from datetime import date
from typing import Optional

from fastapi import Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy import func, select

from .models import Voucher, VoucherLine

STANDARD = "小企业会计准则"

# 资产负债表：(行次, 项目, 取数)；取数 = ("sum", 前缀…) 借方为正 / ("neg", 前缀…) 贷方为正 / ("ar"|"ap"|"adv"|"pre") 往来重分类 / ("row", 行…) 合计 / ("up",) 未分配利润
BS_ASSETS = [
    (1, "货币资金", ("sum", "1001", "1002", "1012")),
    (2, "短期投资", ("sum", "1101")),
    (3, "应收票据", ("sum", "1121")),
    (4, "应收账款", ("ar",)),
    (5, "预付账款", ("pre",)),
    (6, "应收股利", ("sum", "1131")),
    (7, "应收利息", ("sum", "1132")),
    (8, "其他应收款", ("sum", "1221")),
    (9, "存货", ("sum", "1401", "1402", "1403", "1404", "1405", "1406", "1407", "1408", "1411", "1421", "5001", "5101", "5201", "5401")),
    (10, "其中：原材料", ("sum", "1403")),
    (11, "在产品", ("sum", "5001", "5101")),
    (12, "库存商品", ("sum", "1405")),
    (13, "周转材料", ("sum", "1411")),
    (14, "其他流动资产", ("ota",)),
    (15, "流动资产合计", ("row", 1, 2, 3, 4, 5, 6, 7, 8, 9, 14)),
    (16, "长期债券投资", ("sum", "1501")),
    (17, "长期股权投资", ("sum", "1511")),
    (18, "固定资产原价", ("sum", "1601")),
    (19, "减：累计折旧", ("neg", "1602")),
    (20, "固定资产账面价值", ("diff", 18, 19)),
    (21, "在建工程", ("sum", "1604")),
    (22, "工程物资", ("sum", "1605")),
    (23, "固定资产清理", ("sum", "1606")),
    (24, "生产性生物资产", ("sum", "1621", "1622x")),
    (25, "无形资产", ("sum", "1701", "1702x")),
    (26, "开发支出", ("sum", "5301")),
    (27, "长期待摊费用", ("sum", "1801")),
    (28, "其他非流动资产", ("sum", "1901x")),
    (29, "非流动资产合计", ("row", 16, 17, 20, 21, 22, 23, 24, 25, 26, 27, 28)),
    (30, "资产总计", ("row", 15, 29)),
]
BS_LIAB = [
    (31, "短期借款", ("neg", "2001")),
    (32, "应付票据", ("neg", "2201")),
    (33, "应付账款", ("ap",)),
    (34, "预收账款", ("adv",)),
    (35, "应付职工薪酬", ("neg", "2211")),
    (36, "应交税费", ("tax",)),
    (37, "应付利息", ("neg", "2231")),
    (38, "应付利润", ("neg", "2232")),
    (39, "其他应付款", ("neg", "2241")),
    (40, "其他流动负债", ("neg", "2401x")),
    (41, "流动负债合计", ("row", 31, 32, 33, 34, 35, 36, 37, 38, 39, 40)),
    (42, "长期借款", ("neg", "2501")),
    (43, "长期应付款", ("neg", "2701")),
    (44, "递延收益", ("neg", "2401")),
    (45, "其他非流动负债", ("neg", "2901x")),
    (46, "非流动负债合计", ("row", 42, 43, 44, 45)),
    (47, "负债合计", ("row", 41, 46)),
    (48, "实收资本（或股本）", ("neg", "4001")),
    (49, "资本公积", ("neg", "4002")),
    (50, "盈余公积", ("neg", "4101")),
    (51, "未分配利润", ("up",)),
    (52, "所有者权益（或股东权益）合计", ("row", 48, 49, 50, 51)),
    (53, "负债和所有者权益（或股东权益）总计", ("row", 47, 52)),
]
# 利润表：(行次, 项目, 取数, 申报登记键)；费用类借方为正，收入类贷方为正
IS_ROWS = [
    (1, "一、营业收入", ("cr", "6001", "6051"), "revenue"),
    (2, "减：营业成本", ("dr", "6401", "6402"), "cost"),
    (3, "税金及附加", ("dr", "6403"), "taxes"),
    (4, "其中：消费税", ("dr", "640307x"), None),
    (5, "营业税", ("dr", "640308x"), None),
    (6, "城市维护建设税", ("dr", "640301"), None),
    (7, "资源税", ("dr", "640309x"), None),
    (8, "土地增值税", ("dr", "640310x"), None),
    (9, "城镇土地使用税、房产税、车船税、印花税", ("dr", "640304", "640305", "640306"), None),
    (10, "教育费附加、矿产资源补偿费、排污费", ("dr", "640302", "640303"), None),
    (11, "销售费用", ("dr", "6601"), "selling"),
    (12, "其中：商品维修费", ("dr", "660199x"), None),
    (13, "广告费和业务宣传费", ("dr", "660198x"), None),
    (14, "管理费用", ("dr", "6602"), "admin"),
    (15, "其中：开办费", ("dr", "660299x"), None),
    (16, "业务招待费", ("dr", "660298x"), None),
    (17, "研究费用", ("dr", "660211"), None),
    (18, "财务费用", ("dr", "6603"), "finance"),
    (19, "其中：利息费用（收入以“-”号填列）", ("dr", "660302", "660303x"), None),
    (20, "加：投资收益（损失以“-”号填列）", ("cr", "6111"), None),
    (21, "二、营业利润（亏损以“-”号填列）", ("calc", "op"), None),
    (22, "加：营业外收入", ("cr", "6301"), "nonop_income"),
    (23, "其中：政府补助", ("cr", "630101x"), None),
    (24, "减：营业外支出", ("dr", "6711"), "nonop_expense"),
    (25, "其中：坏账损失", ("dr", "671103x"), None),
    (26, "无法收回的长期债券投资损失", ("dr", "671104x"), None),
    (27, "无法收回的长期股权投资损失", ("dr", "671105x"), None),
    (28, "自然灾害等不可抗力因素造成的损失", ("dr", "671106x"), None),
    (29, "税收滞纳金", ("dr", "671107x"), None),
    (30, "三、利润总额（亏损总额以“-”号填列）", ("calc", "total"), "profit_total"),
    (31, "减：所得税费用", ("dr", "6801"), "income_tax"),
    (32, "四、净利润（净亏损以“-”号填列）", ("calc", "net"), "net_profit"),
]
VAT_GROUP = ("222101", "222102", "222103", "222105")   # 销项、进项、已交、转出 → 净借方即留抵
VAT_PENDING = "222104"                                  # 待认证进项
PL_PREFIX = ("6",)   # 损益类（企业会计准则编码 6xxx；表结法：未结转部分并入未分配利润）


def _line_sums(db, start: Optional[date], end: date, by_partner=False, exclude_closing=False):
    q = (select(VoucherLine.account_code, VoucherLine.partner_id, func.sum(VoucherLine.debit - VoucherLine.credit))
         .join(Voucher, VoucherLine.voucher_id == Voucher.id)
         .where(Voucher.status.in_(("posted", "reversed")), Voucher.voucher_date <= end))
    if start:
        q = q.where(Voucher.voucher_date >= start)
    if exclude_closing:
        q = q.where(Voucher.source_type != "closing")
    q = q.group_by(VoucherLine.account_code, VoucherLine.partner_id)
    out = defaultdict(float)
    for code, pid, v in db.execute(q):
        out[(code, pid) if by_partner else code] += v or 0
    return out


def _pick(sums: dict, prefixes) -> float:
    ps = tuple(p for p in prefixes if not p.endswith("x"))     # 以 x 结尾的是占位（本单位暂无该科目）
    return sum(v for code, v in sums.items() if ps and code.startswith(ps))


def balance_sheet(db, asof: date) -> list[dict]:
    sums = _line_sums(db, None, asof)
    by_p = _line_sums(db, None, asof, by_partner=True)
    party = defaultdict(float)             # (类别, 往来单位) → 余额（借正）
    for (code, pid), v in by_p.items():
        for grp, pref in (("1122", "1122"), ("2203", "2203"), ("1123", "1123"), ("2202", "2202"), ("1234", "1234")):
            if code.startswith(pref):
                party[(grp, pid)] += v
    pos = lambda g: sum(v for (k, _), v in party.items() if k == g and v > 0.005)
    neg = lambda g: sum(-v for (k, _), v in party.items() if k == g and v < -0.005)
    vals = {}
    vat = _pick(sums, VAT_GROUP)                       # 应交增值税明细净额：借正＝留抵
    vat_debit = max(vat, 0.0) + max(_pick(sums, (VAT_PENDING,)), 0.0)

    def calc(row, spec):
        kind = spec[0]
        if kind == "ota":
            return _pick(sums, ("1901",)) + vat_debit
        if kind == "tax":
            return -(_pick(sums, ("2221",)) - vat_debit)
        if kind == "sum":
            return _pick(sums, spec[1:])
        if kind == "neg":
            return -_pick(sums, spec[1:])
        if kind == "ar":
            return pos("1122") + pos("2203")
        if kind == "adv":
            return neg("2203") + neg("1122")
        if kind == "pre":
            return pos("1123") + pos("2202") + pos("1234")
        if kind == "ap":
            return neg("2202") + neg("1123") + neg("1234")
        if kind == "row":
            return sum(vals[r] for r in spec[1:])
        if kind == "diff":
            return vals[spec[1]] - vals[spec[2]]
        if kind == "up":
            eq = -_pick(sums, ("4103", "4104", "4199"))
            pl = -sum(v for code, v in sums.items() if code.startswith(PL_PREFIX))
            return eq + pl
        raise ValueError(kind)

    rows = []
    for side, items in (("asset", BS_ASSETS), ("liab", BS_LIAB)):
        for r, label, spec in items:
            vals[r] = round(calc(r, spec), 2) + 0.0
            rows.append({"row": r, "label": label, "side": side, "amount": vals[r],
                         "total": spec[0] in ("row", "diff")})
    return rows


def income_statement(db, start: date, end: date) -> dict:
    sums = _line_sums(db, start, end, exclude_closing=True)
    vals = {}
    for r, label, spec, _ in IS_ROWS:
        kind = spec[0]
        if kind == "dr":
            v = _pick(sums, spec[1:])
        elif kind == "cr":
            v = -_pick(sums, spec[1:])
        elif spec[1] == "op":
            v = vals[1] - vals[2] - vals[3] - vals[11] - vals[14] - vals[18] + vals[20]
        elif spec[1] == "total":
            v = vals[21] + vals[22] - vals[24]
        else:
            v = vals[30] - vals[31]
        vals[r] = round(v, 2) + 0.0
    return vals


def quarter_range(period: str):
    y = int(period[:4])
    if len(period) == 4:
        return date(y, 1, 1), date(y, 12, 31), date(y, 1, 1)
    q = int(period[-1])
    end_m = q * 3
    end = date(y + (end_m == 12), end_m % 12 + 1, 1)
    return date(y, end_m - 2, 1), date.fromordinal(end.toordinal() - 1), date(y, 1, 1)


def statements(db, period: str) -> dict:
    from .tax_filing import TaxFiling, prior_period
    p_start, end, y_start = quarter_range(period)
    ytd = income_statement(db, y_start, end)
    book = income_statement(db, p_start, end)
    prev = prior_period(period)
    prev_row = db.scalar(select(TaxFiling).where(TaxFiling.period == prev)) if prev else None
    prev_book = income_statement(db, y_start, date.fromordinal(p_start.toordinal() - 1)) if prev else {}
    parent_of = {4: 3, 5: 3, 6: 3, 7: 3, 8: 3, 9: 3, 10: 3, 12: 11, 13: 11, 15: 14, 16: 14, 17: 14, 19: 18,
                 23: 22, 25: 24, 26: 24, 27: 24, 28: 24, 29: 24}
    keys = {r: k for r, _, _, k in IS_ROWS}
    filed = (prev_row.data or {}) if prev_row else {}
    filed_rows = filed.get("is") or {}
    is_rows = []
    for r, label, spec, key in IS_ROWS:
        if not prev:
            filed_basis = ytd[r]
        elif not prev_row:
            filed_basis = ytd[r] - prev_book.get(r, 0)          # 上期未登记：按账面上期累计
        elif key:
            filed_basis = ytd[r] - filed.get(key, 0.0)            # 申报表上没有的项目 = 申报为 0
        elif str(r) in filed_rows:
            filed_basis = ytd[r] - (filed_rows[str(r)][1] or 0.0)  # 导入的已申报全表（v0.34）：逐行用申报累计
        else:
            # 「其中」行没有登记申报数：所属大项上期申报为 0 则视同 0，否则按账面上期累计
            pk = keys.get(parent_of.get(r))
            filed_basis = ytd[r] - (0.0 if abs(filed.get(pk, 0.0)) < 0.005 else prev_book.get(r, 0))
        is_rows.append({"row": r, "label": label, "ytd": ytd[r], "period_book": book[r],
                        "period_filed": round(filed_basis, 2) + 0.0, "bold": spec[0] == "calc" or key in ("revenue",)})
    # 申报口径：营业利润/利润总额/净利润按调整后的行重算，保证表内勾稽
    f = {x["row"]: x["period_filed"] for x in is_rows}
    f[21] = round(f[1] - f[2] - f[3] - f[11] - f[14] - f[18] + f[20], 2)
    f[30] = round(f[21] + f[22] - f[24], 2)
    f[32] = round(f[30] - f[31], 2)
    for x in is_rows:
        x["period_filed"] = f[x["row"]] + 0.0
    bs_end = balance_sheet(db, end)
    bs_begin = balance_sheet(db, date(y_start.year - 1, 12, 31))
    begin = {x["row"]: x["amount"] for x in bs_begin}
    # 年初数照抄本年已申报的（用户 2026-10-01 定）：取本年最近一期导入了全表的申报；没有才用 ERP 上年末
    begin_source = "erp"
    for t in db.scalars(select(TaxFiling).where(TaxFiling.period.like(f"{y_start.year}-Q%")).order_by(TaxFiling.period.desc())):
        rows = (t.data or {}).get("bs")
        if rows and t.period < (period if "-Q" in period else f"{period}-Q5"):
            begin = {int(k): v[1] for k, v in rows.items()} | {r: v for r, v in begin.items() if str(r) not in rows}
            begin_source = t.period
            break
    for x in bs_end:
        x["begin"] = begin[x["row"]]
    b = {x["row"]: x["amount"] for x in bs_end}
    from .company import get as company
    return {"standard": company(db).get("accounting_standard") or STANDARD, "company": company(db).get("name", ""), "period": period, "start": str(p_start), "end": str(end), "year_start": str(y_start),
            "prior": prev, "prior_filed": bool(prev_row), "begin_source": begin_source, "balance_sheet": bs_end, "income_statement": is_rows,
            "balanced": abs(b[30] - b[53]) < 0.01}


def to_xlsx(data: dict) -> bytes:
    import openpyxl
    from openpyxl.styles import Alignment, Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "资产负债表"
    ws.append(["资产负债表（会小企01表）"])
    ws.append([f"编制单位：{data.get('company', '')}", "", f"{data['end']}", "", "", "单位：元"])
    ws.append(["资产", "行次", "期末余额", "年初余额", "负债和所有者权益", "行次", "期末余额", "年初余额"])
    a = [x for x in data["balance_sheet"] if x["side"] == "asset"]
    l = [x for x in data["balance_sheet"] if x["side"] == "liab"]
    for i in range(max(len(a), len(l))):
        x = a[i] if i < len(a) else None
        y = l[i] if i < len(l) else None
        ws.append([x["label"] if x else "", x["row"] if x else "", x["amount"] if x else "", x["begin"] if x else "",
                   y["label"] if y else "", y["row"] if y else "", y["amount"] if y else "", y["begin"] if y else ""])
    ws2 = wb.create_sheet("利润表")
    ws2.append(["利润表（会小企02表）"])
    ws2.append([f"编制单位：{data.get('company', '')}", f"{data['start']} ~ {data['end']}", "", "", "单位：元"])
    ws2.append(["项目", "行次", "本年累计金额", "本期金额（申报口径）", "本期金额（账面）"])
    for x in data["income_statement"]:
        ws2.append([x["label"], x["row"], x["ytd"], x["period_filed"], x["period_book"]])
    for w in (ws, ws2):
        w["A1"].font = Font(bold=True, size=14)
        for col in "ABCDEFGH":
            w.column_dimensions[col].width = 16
        w.column_dimensions["A"].width = 38
        w.column_dimensions["E"].width = 30 if w is ws else 18
        for row in w.iter_rows(min_row=4):
            for c in row:
                if isinstance(c.value, float):
                    c.number_format = "#,##0.00"
                    c.alignment = Alignment(horizontal="right")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def install(app, get_db):
    import re

    def _check(period):
        if not re.fullmatch(r"\d{4}(-Q[1-4])?", period):
            raise HTTPException(400, "期间格式：2026-Q3 或 2026")

    @app.get("/api/statements/small")
    def small(period: str, db=Depends(get_db)):
        _check(period)
        return statements(db, period)

    @app.get("/api/statements/small.xlsx")
    def small_xlsx(period: str, db=Depends(get_db)):
        _check(period)
        data = to_xlsx(statements(db, period))
        from urllib.parse import quote
        name = quote(f"小企业会计准则报表_{period}.xlsx")
        return StreamingResponse(io.BytesIO(data), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                 headers={"Content-Disposition": f"attachment; filename*=utf-8''{name}"})
