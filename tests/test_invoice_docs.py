"""发票档案：XML/PDF 文本解析、上传归档去重、关联 ERP 发票、勾选导入 → 待认证转进项、进项转出、增值税台账、原件下载"""
import base64
import io

OWN = "91320000MA1EXAMPL1"


def _xml(no, seller="测试供应商", seller_tax="91110000000000001X", buyer="示例智能科技有限公司", buyer_tax=OWN,
         net="100.00", tax="13.00", total="113.00", kind="增值税专用发票"):
    return (f'<?xml version="1.0" encoding="utf-8"?><EInvoice><Header><EIid>{no}</EIid><InherentLabel><InIssuType><LabelCode>Y</LabelCode></InIssuType>'
            f'<GeneralOrSpecialVAT><LabelName>{kind}</LabelName></GeneralOrSpecialVAT></InherentLabel></Header><EInvoiceData>'
            f'<SellerInformation><SellerIdNum>{seller_tax}</SellerIdNum><SellerName>{seller}</SellerName></SellerInformation>'
            f'<BuyerInformation><BuyerIdNum>{buyer_tax}</BuyerIdNum><BuyerName>{buyer}</BuyerName></BuyerInformation>'
            f'<BasicInformation><TotalAmWithoutTax>{net}</TotalAmWithoutTax><TotalTaxAm>{tax}</TotalTaxAm>'
            f'<TotalTax-includedAmount>{total}</TotalTax-includedAmount><RequestTime>2026-07-05 10:00:00</RequestTime></BasicInformation>'
            f'<IssuItemInformation><ItemName>*电子元件*驱动器</ItemName><TaxRate>0.13</TaxRate></IssuItemInformation></EInvoiceData></EInvoice>').encode()


def _up(client, name, data, direction="in"):
    r = client.post("/api/invoice-docs/upload", json={"name": name, "content_b64": base64.b64encode(data).decode(),
                                                     "direction": direction})
    assert r.status_code == 200, r.text
    return r.json()


def _deduct_xlsx(no, period="202607", tax=13.0):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "发票"
    ws.append(["发票清单"])
    ws.append(["纳税人识别号", OWN, "税款所属期", period])
    ws.append(["序号", "勾选状态", "发票来源", "转内销证明编号", "数电发票号码", "发票代码", "发票号码", "开票日期",
               "销售方纳税人识别号", "销售方纳税人名称", "金额", "税额", "有效抵扣税额", "票种", "票种标签", "发票状态", "勾选时间"])
    ws.append(["1", "已勾选", "电子发票服务平台", None, no, None, None, "2026-07-05", "91110000000000001X", "测试供应商",
               100, tax, tax, "数电发票(增值税专用发票)", "增值税专用发票", "正常", "2026-08-10 10:00:00"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_pdf_text_parser_layout_independent():
    from app.invoice_docs import parse_pdf_text
    text = "\n".join(["示例速运有限公司", "2026年05月21日", "91440300MA5FAKE01X", " 26957000000135531541", OWN,
                      "示例智能科技有限公司", "电子发票（普通发票）", "873.12 6%1 873.12 52.39*快递服务*收派服务费",
                      "¥873.12 ¥52.39", "¥925.51玖佰贰拾伍元伍角壹分"])
    r = parse_pdf_text(text, "in", {"name": "示例智能科技有限公司", "tax": OWN, "keys": ["示例智能科技有限公司", "示例智能"]})
    assert r["invoice_no"] == "26957000000135531541" and r["issue_date"] == "2026-05-21"
    assert r["seller_name"] == "示例速运有限公司" and r["seller_tax_no"] == "91440300MA5FAKE01X"
    assert (r["amount_ex_tax"], r["tax_amount"], r["amount_tax"]) == (873.12, 52.39, 925.51) and r["kind"] == "数电普票"


def test_upload_link_deduct_and_vat_ledger(client, seeded):
    no = "26322000009999990001"
    client.put("/api/partners/1", json={"code": "T01", "name": "测试供应商", "is_supplier": True, "tax_no": "91110000000000001X"})
    inv = client.post("/api/invoices", json={"direction": "input", "invoice_no": no, "invoice_date": "2026-07-05", "partner_id": 1,
                                             "amount_tax": 113, "amount_ex_tax": 100, "tax_amount": 13}).json()
    res = _up(client, f"dzfp_{no}.xml", _xml(no))
    assert res[0]["ok"] and res[0]["direction"] == "in"
    assert _up(client, f"dzfp_{no}.xml", _xml(no))[0]["new_file"] is False            # 同文件不重复存
    d = client.get("/api/invoice-docs").json()[0]
    assert d["invoice_id"] == inv["id"] and d["category"] == "trade" and d["deduct_status"] == "pending" and d["partner_id"] == 1
    assert client.get(f"/api/invoice-docs/files/{d['files'][0]['id']}").status_code == 200
    # 未勾选：税额在待认证，进项为 0
    led = {m["month"]: m for m in client.get("/api/vat-ledger").json()["months"]}
    assert led["2026-07"]["input"] == 0 and led["2026-07"]["pending_change"] == 13
    # 导入勾选清单 → 所属期 202607 出 待认证→进项 凭证
    r = _up(client, "用途确认信息202607.xlsx", _deduct_xlsx(no))
    assert r[0]["deducted"] == 1
    d = client.get("/api/invoice-docs").json()[0]
    assert d["deduct_status"] == "deducted" and d["deduct_period"] == "2026-07"
    led = {m["month"]: m for m in client.get("/api/vat-ledger").json()["months"]}
    assert led["2026-07"]["input"] == 13 and led["2026-07"]["carry_out"] == 13
    # 进项转出（如集体福利）
    r = client.put(f"/api/invoice-docs/{d['id']}", json={"deduct_status": "transferred", "transfer_period": "2026-09",
                                                         "transfer_account": "660208"})
    assert r.status_code == 200, r.text
    led = {m["month"]: m for m in client.get("/api/vat-ledger").json()["months"]}
    assert led["2026-09"]["transfer_out"] == 13 and led["2026-09"]["carry_in"] == 13 and led["2026-09"]["payable"] == 0
    assert client.get("/api/change-log?table=invoice_doc").json()[0]["action"] == "修改发票档案"


def test_ofd_needs_existing_doc_and_voucher_link(client, seeded):
    no = "26322000009999990002"
    r = _up(client, f"dzfp_{no}_某公司.ofd", b"PK-not-really")
    assert r[0]["ok"] is False
    _up(client, f"{no}.xml", _xml(no, kind="普通发票"))
    assert _up(client, f"dzfp_{no}_某公司.ofd", b"PK-not-really")[0]["ok"]
    d = client.get("/api/invoice-docs").json()[0]
    assert len(d["files"]) == 2 and d["deduct_status"] == "nondeductible" and not d["booked"]
    v = client.post("/api/vouchers", json={"voucher_date": "2026-07-05", "summary": "费用", "lines": [
        {"account": "660210", "debit": 113}, {"account": "2241", "credit": 113}]}).json()
    assert client.put(f"/api/invoice-docs/{d['id']}", json={"voucher_no": v["voucher_no"], "category": "expense"}).status_code == 200
    d = client.get("/api/invoice-docs").json()[0]
    assert d["booked"] and d["booked_month"] == "2026-07" and d["voucher_no"] == v["voucher_no"]


def test_tax_filing_current_period_is_ytd_minus_prior_filed(client, seeded):
    client.post("/api/vouchers", json={"voucher_date": "2026-02-10", "summary": "q1 费用", "lines": [
        {"account": "660210", "debit": 100}, {"account": "2241", "credit": 100}]})
    client.post("/api/vouchers", json={"voucher_date": "2026-08-10", "summary": "q3 费用", "lines": [
        {"account": "660201", "debit": 50}, {"account": "2241", "credit": 50}]})
    # Q1 实际申报管理费用 0、利润 0（与 ERP 的 100 不同）→ 差异落到 Q2
    assert client.put("/api/tax-filings", json={"period": "2026-Q1", "data": {"admin": 0, "profit_total": 0}}).status_code == 200
    r = client.get("/api/tax-filings/report?period=2026-Q2").json()
    rows = {x["key"]: x for x in r["rows"]}
    assert rows["admin"]["erp_ytd"] == 100 and rows["admin"]["current"] == 100 and rows["profit_total"]["current"] == -100
    r = client.get("/api/tax-filings/report?period=2026-Q3").json()
    assert r["prior_missing"] is True
    client.put("/api/tax-filings", json={"period": "2026-Q2", "data": {"admin": 100, "profit_total": -100}})
    rows = {x["key"]: x for x in client.get("/api/tax-filings/report?period=2026-Q3").json()["rows"]}
    assert rows["admin"]["erp_ytd"] == 150 and rows["admin"]["current"] == 50


def test_small_enterprise_statements(client, seeded):
    # 股东出资、预付/应付往来重分类、费用；待认证进项列其他流动资产，不与欠税局的销项相抵（财会〔2016〕22 号）
    for code, name in (("400101", "股东甲"),):
        client.post("/api/accounts", json={"code": code, "name": name, "acc_type": "equity", "direction": "C", "parent_code": "4001"})
    client.post("/api/vouchers", json={"voucher_date": "2026-01-05", "summary": "出资", "lines": [
        {"account": "100201", "debit": 1000}, {"account": "400101", "credit": 1000}]})
    client.post("/api/vouchers", json={"voucher_date": "2026-08-05", "summary": "费用含进项", "lines": [
        {"account": "660210", "debit": 100}, {"account": "222104", "debit": 13}, {"account": "100201", "credit": 113}]})
    # 同一单位：应付贷方 50（应列应付账款）；另一单位预付借方 30（应列预付账款）
    client.post("/api/vouchers", json={"voucher_date": "2026-08-06", "summary": "往来", "lines": [
        {"account": "660210", "debit": 50}, {"account": "2202", "credit": 50, "partner_id": 1}]})
    client.post("/api/vouchers", json={"voucher_date": "2026-08-07", "summary": "预付", "lines": [
        {"account": "1123", "debit": 30, "partner_id": 2}, {"account": "100201", "credit": 30}]})
    client.post("/api/vouchers", json={"voucher_date": "2026-08-08", "summary": "销项", "lines": [
        {"account": "100201", "debit": 20}, {"account": "222101", "credit": 20}]})
    client.put("/api/tax-filings", json={"period": "2026-Q2", "data": {"admin": 0, "profit_total": 0}})
    d = client.get("/api/statements/small?period=2026-Q3").json()
    bs = {x["row"]: x["amount"] for x in d["balance_sheet"]}
    assert d["balanced"] and bs[1] == 877 and bs[5] == 30 and bs[33] == 50 and bs[14] == 13 and bs[36] == 20
    assert bs[48] == 1000 and bs[51] == -150 and bs[30] == bs[53] == 920
    tb = {x["key"]: x["erp"] for x in client.get("/api/tax-filings/report?period=2026-Q3").json()["balance"]}
    assert tb["total_assets"] == 920 and tb["total_liabilities"] == 70
    isr = {x["row"]: x for x in d["income_statement"]}
    assert isr[14]["ytd"] == 150 and isr[14]["period_filed"] == 150 and isr[32]["period_filed"] == -150
    r = client.get("/api/statements/small.xlsx?period=2026-Q3")
    assert r.status_code == 200 and r.content[:2] == b"PK"


def test_register_trade_invoice_from_archive(client, seeded):
    no = "26312000009999990003"
    client.put("/api/partners/1", json={"code": "T01", "name": "测试供应商", "is_supplier": True, "tax_no": "91110000000000001X"})
    c = client.post("/api/contracts", json={"contract_no": "P-REG-1", "contract_type": "purchase", "partner_id": 1,
                                            "sign_date": "2026-07-01", "amount_tax": 113, "lines": [
                                                {"material_id": seeded["mat1"], "qty": 1, "price_tax": 113}]}).json()
    _up(client, f"{no}.xml", _xml(no))
    d = client.get("/api/invoice-docs").json()[0]
    assert d["invoice_id"] is None and d["partner_id"] == 1
    r = client.post(f"/api/invoice-docs/{d['id']}/register", json={"contract_id": c["id"]})
    assert r.status_code == 200, r.text
    d = client.get("/api/invoice-docs").json()[0]
    assert d["invoice_id"] and d["category"] == "trade" and d["booked_month"] == "2026-07"
    assert client.get(f"/api/contracts/{c['id']}").json()["invoice_sum"] == 113
    assert client.post(f"/api/invoice-docs/{d['id']}/register", json={}).status_code == 400


def test_close_check(client, seeded):
    r = client.get("/api/close-check?month=2026-09")
    assert r.status_code == 200, r.text
    d = r.json()
    keys = {x["key"]: x for x in d["items"]}
    assert {"bank_coverage", "bank_unmatched", "invoice_unbooked", "vat_deduct", "allocation", "payroll",
            "depreciation", "balanced", "lock"} <= set(keys)
    assert keys["lock"]["status"] == "todo" and d["locked"] is False
    # 月结检查只读：自动核销试算不落库
    before = client.get("/api/allocations").json()
    client.get("/api/close-check?month=2026-09")
    assert client.get("/api/allocations").json() == before


def test_stamp_contract_pdf(client, seeded, tmp_path):
    import pymupdf
    # 造一张带透明通道的章图 + 一份有「甲方（盖章）」字样的合同 PDF（采购合同我方是甲方）
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 200, 200), 1)
    pix.clear_with(0)
    png = pix.tobytes("png")
    r = client.post("/api/seals", json={"name": "合同章", "content_b64": base64.b64encode(png).decode(), "dpi": 125})
    assert r.status_code == 200, r.text
    assert abs(r.json()["width_mm"] - 40.64) < 0.01          # 200 px / 125 dpi = 1.6 in
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_font(fontname="china", fontbuffer=pymupdf.Font("china-s").buffer)
    page.insert_text((72, 700), "甲方（盖章）：", fontname="china", fontsize=12)
    pdf = doc.tobytes()
    c = client.post("/api/contracts", json={"contract_no": "P-STAMP-1", "contract_type": "purchase", "partner_id": 1,
                                            "sign_date": "2026-09-01", "amount_tax": 113, "lines": [
                                                {"material_id": seeded["mat1"], "qty": 1, "price_tax": 113}]}).json()
    f = client.post(f"/api/contracts/{c['id']}/files", json={"name": "合同.pdf", "content_b64": base64.b64encode(pdf).decode(),
                                                            "kind": "other"}).json()
    pv = client.post(f"/api/contract-files/{f['id']}/stamp-preview", json={"seal": "合同章", "mode": "anchor"})
    assert pv.status_code == 200, pv.text
    assert pv.json()["anchor"] == "甲方（盖章）" and pv.json()["png_b64"]
    st = client.post(f"/api/contract-files/{f['id']}/stamp", json={"seal": "合同章", "mode": "anchor"}).json()
    assert st["kind"] == "other" and st["file_name"].endswith("-我方合同章.pdf")
    files = client.get(f"/api/contracts/{c['id']}/files").json()
    assert len(files) == 2 and "待对方回签" in [x for x in files if x["id"] == st["id"]][0]["note"]
    bad = client.post(f"/api/contract-files/{f['id']}/stamp", json={"seal": "合同章", "mode": "mirror"})
    assert bad.status_code == 400 and "对方红章" in bad.json()["detail"]


def test_invoice_export_template(client, seeded):
    import io as _io
    import openpyxl
    client.put("/api/partners/2", json={"code": "T02", "name": "测试客户", "is_customer": True, "tax_no": "91440300MA5FAKE02X",
                                        "address": "深圳某地", "phone": "0755-1", "bank_name": "某银行", "bank_account": "6222 0001"})
    c = client.post("/api/contracts", json={"contract_no": "XS2699001-新", "contract_type": "sales", "partner_id": 2,
                                            "sign_date": "2026-09-01", "amount_tax": 1130, "lines": [
                                                {"material_id": seeded["mat1"], "qty": 10, "price_tax": 113}]}).json()
    pv = client.get(f"/api/contracts/{c['id']}/invoice-export").json()
    assert pv["serial"] == "XS2699001-" and pv["lines"][0]["tax_code"] == "1090131050000000000"
    r = client.post(f"/api/contracts/{c['id']}/invoice-export", json={"serial": pv["serial"], "lines": [
        {"line_id": pv["lines"][0]["line_id"], "qty": 4}]})
    assert r.status_code == 200, r.text
    wb = openpyxl.load_workbook(_io.BytesIO(r.content))
    s1, s2 = wb["1-发票基本信息"], wb["2-发票明细信息"]
    h1 = {s1.cell(3, i).value: s1.cell(4, i).value for i in range(1, s1.max_column + 1)}
    assert h1["发票类型"] == "增值税专用发票" and h1["购买方纳税人识别号"] == "91440300MA5FAKE02X" and h1["是否含税"] == "是"
    assert h1["购买方银行账号"] == "62220001" and h1["备注"] == "合同XS2699001-新"
    h2 = {s2.cell(3, i).value: s2.cell(4, i).value for i in range(1, s2.max_column + 1)}
    assert h2["数量"] == "4" and h2["单价"] == "113" and h2["金额"] == "452.00" and h2["税率"] == "0.13"
    assert "excelVersion" in wb.sheetnames


def test_tax_form_export_import_roundtrip(client, seeded):
    """税局财务报表导入文件：导出 → （网页上改一个数）→ 导回登记；下季本期按申报累计扣减，年初照抄已申报"""
    import base64, io, openpyxl
    client.post("/api/vouchers", json={"voucher_date": "2026-02-10", "summary": "费用", "lines": [
        {"account": "660204", "debit": 100}, {"account": "100201", "credit": 100}]})
    r = client.get("/api/tax-filings/export.xlsx?period=2026-Q1")
    assert r.status_code == 200, r.text
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    bs, is_ = wb["资产负债表"], wb["利润表_月季报"]
    assert bs["D4"].value == "2026-01-01" and bs["H4"].value == "2026-03-31" and is_["D4"].value == "2026-01-01"
    assert not any(isinstance(c.value, str) and c.value.startswith("=") for ws in wb for row in ws.iter_rows() for c in row)
    row14 = next(c.row for c in is_["C"] if c.value == 14)
    assert is_.cell(row14, 4).value == 100 and is_.cell(row14, 5).value == 100
    is_.cell(row14, 4).value = is_.cell(row14, 5).value = 90          # 用户在税局网页上改成 90
    row1 = next(c.row for c in bs["C"] if c.value == 1)
    bs.cell(row1, 5).value = 12345                                     # 年初货币资金（已申报）
    buf = io.BytesIO(); wb.save(buf)
    res = client.post("/api/tax-filings/import", json={"name": "q1.xlsx", "content_b64": base64.b64encode(buf.getvalue()).decode()})
    assert res.status_code == 200, res.text
    j = res.json()
    assert j["period"] == "2026-Q1"
    assert {(d["row"], d["col"], d["filed"]) for d in j["diffs"] if d["sheet"] == "利润表" and d["row"] == 14} == {(14, "本期", 90), (14, "本年累计", 90)}
    t = next(x for x in client.get("/api/tax-filings").json() if x["period"] == "2026-Q1")
    assert t["data"]["admin"] == 90 and t["data"]["files"] == ["2026-Q1/q1.xlsx"]
    assert client.get("/api/tax-filings/file?path=2026-Q1/q1.xlsx").status_code == 200
    assert client.get("/api/tax-filings/file?path=../owe.db").status_code == 404
    # 手工登记不冲掉导入的全表
    client.put("/api/tax-filings", json={"period": "2026-Q1", "data": {**{k: v for k, v in t["data"].items() if isinstance(v, (int, float))}}})
    t = next(x for x in client.get("/api/tax-filings").json() if x["period"] == "2026-Q1")
    assert "bs" in t["data"] and t["data"]["files"]
    # Q2：本期 = ERP 累计 100 − 申报累计 90 = 10；年初照抄 Q1 申报
    st = client.get("/api/statements/small?period=2026-Q2").json()
    assert next(x for x in st["income_statement"] if x["row"] == 14)["period_filed"] == 10
    assert st["begin_source"] == "2026-Q1" and next(x for x in st["balance_sheet"] if x["row"] == 1)["begin"] == 12345
    assert client.get("/api/tax-filings/export.xlsx?period=2026").status_code == 400


def test_tax_report_payroll_and_opening_assets(client, seeded):
    v = lambda d, lines: client.post("/api/vouchers", json={"voucher_date": d, "summary": "x", "lines": lines})
    v("2026-07-31", [{"account": "660201", "debit": 8500}, {"account": "221101", "credit": 8500},
                     {"account": "221101", "debit": 1917.5}, {"account": "221102", "credit": 1917.5}])
    v("2026-08-10", [{"account": "221101", "debit": 6582.5}, {"account": "100201", "credit": 6582.5}])
    v("2026-08-20", [{"account": "221102", "debit": 1917.5}, {"account": "100201", "credit": 1917.5}])
    client.put("/api/tax-filings", json={"period": "2026-Q2", "data": {"total_assets": 8568784.04}})
    r = client.get("/api/tax-filings/report?period=2026-Q3").json()
    assert r["payroll"] == {"expensed": 8500, "paid": 8500}           # 6582.5 + 1917.5，内部代扣分录不重复算
    assert r["assets_open"] == 8568784.04
